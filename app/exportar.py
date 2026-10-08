"""Exportación del listado a Excel (.xlsx) y CSV."""
from __future__ import annotations

import csv
from pathlib import Path

from .modelo import Cliente, Factura

COLUMNAS = ["Cliente", "NIF/CIF", "Tipo", "Nº factura", "Fecha", "Base imponible", "% IVA",
            "Cuota IVA", "% R.E.", "Recargo eq.", "% IRPF", "IRPF", "Total factura"]


def _filas(facturas: list[Factura], clientes: dict[str, Cliente]):
    for f in facturas:
        cli = clientes.get(f.cliente_clave)
        lineas = f.lineas or [None]
        for i, l in enumerate(lineas):
            primero = i == 0
            yield [
                cli.nombre if cli else "", cli.nif if cli else "", f.tipo, f.numero, f.fecha,
                l.base if l else 0.0, l.iva_pct if l else 0.0, l.iva if l else 0.0,
                l.re_pct if l else 0.0, l.re if l else 0.0,
                f.irpf_pct if primero else 0.0, f.irpf if primero else 0.0,
                f.importe_total if primero else None,
            ]


def exportar_excel(ruta, facturas: list[Factura], clientes: dict[str, Cliente], empresa: dict,
                   titulo: str = "Listado de facturas") -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    ruta = Path(ruta)
    wb = Workbook()
    ws = wb.active
    ws.title = "Facturas"
    acento = (empresa.get("color") or "#C2410C").lstrip("#").upper()

    ws["A1"] = empresa.get("nombre", "")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = titulo
    ws["A2"].font = Font(bold=True, size=11, color=acento)
    fila0 = 4
    for c, nombre in enumerate(COLUMNAS, 1):
        cel = ws.cell(row=fila0, column=c, value=nombre)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="1F2328")
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila0].height = 28

    r = fila0
    for fila in _filas(facturas, clientes):
        r += 1
        for c, v in enumerate(fila, 1):
            cel = ws.cell(row=r, column=c, value=v)
            if c == 5 and v:
                cel.number_format = "DD/MM/YYYY"
            elif c in (7, 9, 11):
                cel.number_format = '0.00" %";-0.00" %";"-"'
            elif c >= 6:
                cel.number_format = '#,##0.00 €;-#,##0.00 €;"-"'
    ultima = r
    r += 1
    ws.cell(row=r, column=1, value="TOTALES").font = Font(bold=True)
    for c in (6, 8, 10, 12, 13):
        col = get_column_letter(c)
        cel = ws.cell(row=r, column=c, value=f"=SUM({col}{fila0 + 1}:{col}{ultima})")
        cel.font = Font(bold=True)
        cel.number_format = '#,##0.00 €'
    linea = Side(style="medium", color="1F2328")
    for c in range(1, len(COLUMNAS) + 1):
        ws.cell(row=r, column=c).border = Border(top=linea)

    anchos = [34, 13, 13, 15, 11, 14, 8, 12, 8, 12, 8, 11, 14]
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.freeze_panes = ws.cell(row=fila0 + 1, column=1)
    ws.auto_filter.ref = f"A{fila0}:{get_column_letter(len(COLUMNAS))}{ultima}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta


def exportar_csv(ruta, facturas: list[Factura], clientes: dict[str, Cliente]) -> Path:
    """CSV con ';' y coma decimal, que Excel en español abre directamente."""
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)

    def fmt(v):
        if v is None:
            return ""
        if isinstance(v, float):
            return f"{v:.2f}".replace(".", ",")
        if hasattr(v, "strftime"):
            return v.strftime("%d/%m/%Y")
        return str(v)

    with ruta.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(COLUMNAS)
        for fila in _filas(facturas, clientes):
            w.writerow([fmt(v) for v in fila])
    return ruta
